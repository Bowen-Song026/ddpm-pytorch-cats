import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = load_workbook("d:/denoising-diffusion-pytorch-main/denoising-diffusion-pytorch-main/实验结果对比表.xlsx")

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
center_alignment = Alignment(horizontal="center", vertical="center")
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

if "学习率实验" in wb.sheetnames:
    wb.remove(wb["学习率实验"])
ws = wb.create_sheet(title="学习率实验")

lr_data = [
    ["实验名称", "学习率", "dim", "batch_size", "grad_accum", "有效batch", "训练步数", "速度(steps/s)", "总时间(分)", "最终Loss", "Loss STD", "学习效果", "图片位置"],
    ["lr_1e-5", 1e-5, 64, 4, 4, 16, 2000, 1.96, 17.1, 0.0736, 0.1317, "学习太慢", "./lr_experiment_lr_1e-5/sample.png"],
    ["lr_4e-5", 4e-5, 64, 4, 4, 16, 2000, 1.86, 17.9, 0.0547, 0.1825, "学习正常", "./lr_experiment_lr_4e-5/sample.png"],
    ["lr_8e-5", 8e-5, 64, 4, 4, 16, 2000, 1.80, 18.5, 0.0509, 0.1096, "学习较好", "./lr_experiment_lr_8e-5/sample.png"],
    ["lr_1e-4", 1e-4, 64, 4, 4, 16, 2000, 1.93, 17.3, 0.0437, 0.1278, "最优", "./lr_experiment_lr_1e-4/sample.png"],
    ["lr_2e-4", 2e-4, 64, 4, 4, 16, 2000, 1.78, 18.7, 0.0560, 0.0919, "学习太快", "./lr_experiment_lr_2e-4/sample.png"],
]

for r_idx, row in enumerate(lr_data, 1):
    for c_idx, value in enumerate(row, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=value)
        cell.alignment = center_alignment
        cell.border = thin_border
        if r_idx == 1:
            cell.font = header_font
            cell.fill = header_fill

for col in ws.columns:
    max_length = 0
    column = col[0].column_letter
    for cell in col:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = max(max_length + 2, 12)
    ws.column_dimensions[column].width = adjusted_width

wb.save("d:/denoising-diffusion-pytorch-main/denoising-diffusion-pytorch-main/实验结果对比表.xlsx")
print("Excel文件已更新，添加了学习率实验数据！")