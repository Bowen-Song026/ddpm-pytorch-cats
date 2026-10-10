import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = load_workbook("d:/denoising-diffusion-pytorch-main/denoising-diffusion-pytorch-main/实验结果对比表.xlsx")

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
center_alignment = Alignment(horizontal="center", vertical="center")
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

ws = wb["采样步实验"]
ws["H1"] = "图片质量评分(1-5)"
ws["I1"] = "分析说明"
ws["H1"].font = header_font
ws["H1"].fill = header_fill
ws["H1"].alignment = center_alignment
ws["H1"].border = thin_border
ws["I1"].font = header_font
ws["I1"].fill = header_fill
ws["I1"].alignment = center_alignment
ws["I1"].border = thin_border

sampling_scores = [
    ["H2", 2.0, "噪声明显，轮廓模糊"],
    ["H3", 3.0, "能看出猫的轮廓，但细节少"],
    ["H4", 3.5, "较好的平衡，毛发开始清晰"],
    ["H5", 4.0, "更清晰的细节，纹理可见"],
    ["H6", 4.5, "最清晰，噪声最小，细节丰富"],
]

for cell, score, desc in sampling_scores:
    ws[cell] = score
    ws[cell].alignment = center_alignment
    ws[cell].border = thin_border
    ws[cell.replace("H", "I")] = desc
    ws[cell.replace("H", "I")].alignment = center_alignment
    ws[cell.replace("H", "I")].border = thin_border

ws = wb["目标函数实验"]
ws["J1"] = "图片质量评分(1-5)"
ws["K1"] = "分析说明"
ws["J1"].font = header_font
ws["J1"].fill = header_fill
ws["J1"].alignment = center_alignment
ws["J1"].border = thin_border
ws["K1"].font = header_font
ws["K1"].fill = header_fill
ws["K1"].alignment = center_alignment
ws["K1"].border = thin_border

objective_scores = [
    ["J2", 3.0, "训练稳定，能生成基本形状，细节一般"],
    ["J3", 1.0, "训练不稳定，生成质量差，像噪声"],
    ["J4", 3.5, "训练稳定，质量较好，轮廓清晰"],
]

for cell, score, desc in objective_scores:
    ws[cell] = score
    ws[cell].alignment = center_alignment
    ws[cell].border = thin_border
    ws[cell.replace("J", "K")] = desc
    ws[cell.replace("J", "K")].alignment = center_alignment
    ws[cell.replace("J", "K")].border = thin_border

wb.save("d:/denoising-diffusion-pytorch-main/denoising-diffusion-pytorch-main/实验结果对比表.xlsx")
print("Excel文件已更新，添加了图片分析评分！")