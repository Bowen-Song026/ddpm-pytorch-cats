import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os

wb = openpyxl.Workbook()
ws = wb.active
ws.title = '参数对比'

header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_font = Font(color='FFFFFF', bold=True)
center_align = Alignment(horizontal='center', vertical='center')
border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)

data = [
    ['参数', '32×32版 (real_cat_results)', '优化版 (optimized_cat_images)', '混合版 (mixed_cat_results)'],
    ['数据集', 'CIFAR-10 猫类', 'CIFAR-10 猫类', 'Oxford-IIIT Pet + CIFAR-10'],
    ['图片数量', '5000张', '5000张', '5554张'],
    ['图像分辨率', '32×32', '64×64', '64×64'],
    ['模型尺寸 (dim)', '64', '64', '64'],
    ['dim_mults', '(1, 2, 4, 8)', '(1, 2, 4, 8)', '(1, 2, 4, 8)'],
    ['Flash Attention', 'False', 'False', 'False'],
    ['时间步 (timesteps)', '1000', '1000', '1000'],
    ['采样步 (sampling_timesteps)', '250', '50', '50'],
    ['预测目标 (objective)', 'pred_noise', 'pred_v', 'pred_v'],
    ['训练步数', '10000', '20000', '20000'],
    ['batch_size', '8', '4', '4'],
    ['梯度累积', '2 (有效batch=16)', '4 (有效batch=16)', '4 (有效batch=16)'],
    ['学习率', '8e-5', '8e-5', '8e-5'],
    ['EMA衰减', '0.995', '0.995', '0.995'],
    ['混合精度 (AMP)', 'True', 'True', 'True'],
    ['GPU', 'RTX 3050 4GB', 'RTX 3050 4GB', 'RTX 3050 4GB'],
]

for row_idx, row_data in enumerate(data, 1):
    for col_idx, cell_value in enumerate(row_data, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=cell_value)
        cell.border = border
        cell.alignment = center_align
        if row_idx == 1:
            cell.fill = header_fill
            cell.font = header_font

ws.column_dimensions['A'].width = 22
ws.column_dimensions['B'].width = 35
ws.column_dimensions['C'].width = 40
ws.column_dimensions['D'].width = 40

output_path = '参数对比表.xlsx'
wb.save(output_path)
print(f"Excel文件已生成: {output_path}")

if os.path.exists(output_path):
    os.startfile(output_path)
    print("Excel已打开")