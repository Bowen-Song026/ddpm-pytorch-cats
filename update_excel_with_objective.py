import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

wb = load_workbook("d:/denoising-diffusion-pytorch-main/denoising-diffusion-pytorch-main/实验结果对比表.xlsx")

header_font = Font(bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
center_alignment = Alignment(horizontal="center", vertical="center")
thin_border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

if "目标函数实验" in wb.sheetnames:
    wb.remove(wb["目标函数实验"])
ws = wb.create_sheet(title="目标函数实验")

objective_data = [
    ["实验名称", "目标函数", "训练步数", "速度(steps/s)", "总时间(分)", "最终Loss", "Loss STD", "训练稳定性", "图片位置"],
    ["obj_pred_noise", "pred_noise", 3000, 2.42, 20.7, 0.0161, 0.1534, "中等", "./objective_experiment_obj_pred_noise/sample.png"],
    ["obj_pred_x0", "pred_x0", 3000, 2.32, 21.5, 0.0825, 20.4854, "差", "./objective_experiment_obj_pred_x0/sample.png"],
    ["obj_pred_v", "pred_v", 3000, 2.33, 21.5, 0.0380, 0.1164, "好", "./objective_experiment_obj_pred_v/sample.png"],
]

for r_idx, row in enumerate(objective_data, 1):
    for c_idx, value in enumerate(row, 1):
        cell = ws.cell(row=r_idx, column=c_idx, value=value)
        cell.alignment = center_alignment
        cell.border = thin_border
        if r_idx == 1:
            cell.font = header_font
            cell.fill = header_fill

for col in ws.columns:
    max_length = 0
    column = col[0].column_letter
    for cell in col:
        try:
            if len(str(cell.value)) > max_length:
                max_length = len(str(cell.value))
        except:
            pass
    adjusted_width = max(max_length + 2, 12)
    ws.column_dimensions[column].width = adjusted_width

wb.save("d:/denoising-diffusion-pytorch-main/denoising-diffusion-pytorch-main/实验结果对比表.xlsx")
print("Excel文件已更新，添加了目标函数实验数据！")